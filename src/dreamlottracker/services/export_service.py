from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from dreamlottracker.services.property_service import PropertyService


class ExportService:
    def __init__(self):
        self.property_service = PropertyService()

    def export_properties_to_excel(self, output_path: Path) -> Path:
        properties = self.property_service.get_all_properties()

        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Properties"

        headers = [
            "Address",
            "City",
            "County",
            "State",
            "Price",
            "Acres",
            "Price Per Acre",
            "Status",
            "Dream Score",
            "Recommendation",
        ]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)

        for row_number, property_ in enumerate(properties, start=2):
            listing = property_.listings[0] if property_.listings else None
            scores = property_.scores

            price = listing.asking_price if listing else 0
            status = listing.status if listing else "Unknown"
            dream_score = scores.dream_score if scores and scores.dream_score is not None else 0
            recommendation = scores.recommendation if scores and scores.recommendation else ""

            sheet.cell(row=row_number, column=1, value=property_.address)
            sheet.cell(row=row_number, column=2, value=property_.city)
            sheet.cell(row=row_number, column=3, value=property_.county)
            sheet.cell(row=row_number, column=4, value=property_.state)
            sheet.cell(row=row_number, column=5, value=price)
            sheet.cell(row=row_number, column=6, value=property_.acres)
            sheet.cell(row=row_number, column=7, value=f'=IF(F{row_number}>0,E{row_number}/F{row_number},"")')
            sheet.cell(row=row_number, column=8, value=status)
            sheet.cell(row=row_number, column=9, value=dream_score)
            sheet.cell(row=row_number, column=10, value=recommendation)

        if len(properties) > 0:
            table_ref = f"A1:J{len(properties) + 1}"
            table = Table(displayName="PropertiesTable", ref=table_ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            sheet.add_table(table)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for column_cells in sheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                value = cell.value
                if value is not None:
                    max_length = max(max_length, len(str(value)))

            sheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

        workbook.save(output_path)
        return output_path
